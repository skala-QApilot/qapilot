"""도메인 지식 Tool — 문서 파싱.

PDF·DOCX 등은 docling, XLSX는 openpyxl로 파싱한다.
파싱 결과를 _chunker 모듈에 위임하여 청크로 분할하고 민감정보를 마스킹한다.

Author: 전아린
Created: 2026-05-07
"""

from __future__ import annotations

import uuid
from pathlib import Path

from qapilot.shared.errors import ErrorCode, ToolExecutionError
from qapilot.tools.domain_knowledge._chunker import mask_sensitive, split_markdown_sections, split_text


class DocumentParser:
    """문서 파일을 파싱하여 청크 목록으로 변환한다.

    역할: PDF·DOCX·XLSX 파싱, _chunker에 청킹·마스킹 위임
    입력: 파일 경로
    출력: list[dict] — chunk_id, source, section, text
    """

    def parse_and_chunk(self, file_path: Path) -> list[dict]:
        """문서를 파싱하고 민감정보 마스킹 후 청크로 분할한다.

        Args:
            file_path: 파싱할 파일 경로.

        Returns:
            list[dict]: chunk_id, source, section, text 키를 가진 청크 목록.

        Raises:
            ToolExecutionError: 문서 파싱에 실패한 경우 (TOOL_005).
        """
        try:
            if file_path.suffix.lower() in (".xlsx", ".xls"):
                blocks = self._parse_excel(file_path)
            elif file_path.suffix.lower() == ".md":
                blocks = self._parse_markdown(file_path)
            else:
                try:
                    blocks = self._parse_with_docling(file_path)
                except Exception:
                    blocks = self._parse_pdf_fallback(file_path)
        except ToolExecutionError:
            raise
        except Exception as e:
            raise ToolExecutionError(
                ErrorCode.TOOL_005,
                f"문서 파싱 실패: {e}. 수동 입력은 add_rule을 사용하세요.",
            ) from e

        chunks = []
        for block in blocks:
            masked = mask_sensitive(block["text"])
            for chunk_text in split_text(masked):
                chunks.append(
                    {
                        "chunk_id": str(uuid.uuid4()),
                        "source": file_path.name,
                        "section": block.get("section", ""),
                        "text": chunk_text,
                    }
                )
        return chunks

    def _parse_with_docling(self, file_path: Path) -> list[dict]:
        """docling으로 PDF·DOCX 등 문서를 파싱하여 섹션 블록 목록을 반환한다.

        Args:
            file_path: 파싱할 파일 경로.

        Returns:
            list[dict]: text, section 키를 가진 섹션 블록 목록.
        """
        from docling.document_converter import DocumentConverter

        result = DocumentConverter().convert(str(file_path))
        markdown = result.document.export_to_markdown()
        return split_markdown_sections(markdown, file_path.stem)

    def _parse_pdf_fallback(self, file_path: Path) -> list[dict]:
        """pypdfium2로 PDF를 페이지 단위 블록으로 파싱하는 폴백 메서드.

        docling 파싱이 실패한 경우(인코딩 오류 등) 호출된다.

        Args:
            file_path: 파싱할 PDF 파일 경로.

        Returns:
            list[dict]: text, section(page_N) 키를 가진 페이지 블록 목록.
        """
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(str(file_path))
        blocks = []
        for i in range(len(pdf)):
            page = pdf[i]
            textpage = page.get_textpage()
            text = textpage.get_text_range()
            if text.strip():
                blocks.append({"text": text, "section": f"page_{i + 1}"})
        return blocks

    @staticmethod
    def _parse_markdown(file_path: Path) -> list[dict]:
        """마크다운 파일을 헤더 기준으로 섹션 블록 목록으로 파싱한다."""
        text = file_path.read_text(encoding="utf-8", errors="replace")
        return split_markdown_sections(text, file_path.stem)

    def _parse_excel(self, file_path: Path) -> list[dict]:
        """openpyxl로 XLSX 파일을 파싱하여 시트별 블록 목록을 반환한다.

        Args:
            file_path: 파싱할 Excel 파일 경로.

        Returns:
            list[dict]: text(시트 전체 행), section(시트명) 키를 가진 블록 목록.
        """
        import openpyxl

        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        blocks = []
        for sheet_name in wb.sheetnames:
            rows = [
                "\t".join(str(c) for c in row if c is not None)
                for row in wb[sheet_name].iter_rows(values_only=True)
            ]
            text = "\n".join(r for r in rows if r.strip())
            if text:
                blocks.append({"text": text, "section": sheet_name})
        wb.close()
        return blocks
